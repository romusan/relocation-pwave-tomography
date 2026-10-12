from pathlib import Path
import csv, hashlib, json, math, shutil, zipfile
from urllib.parse import unquote
from docx import Document

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'articulo_boletin_geologia_uis_adjunto'
RUN=ROOT/'resultados/revision_cientifica_20261011'
DOCX=OUT/'Manuscrito_tomografia_Vp_Colombia_UIS_revision_20261011d.docx'
sha=lambda f:hashlib.sha256(f.read_bytes()).hexdigest()
D=Document(DOCX)
assert len(D.inline_shapes)==7 and len(D.tables)==7
links=[]
for rel in D.part.rels.values():
    if rel.reltype.endswith('/hyperlink') and not rel.target_ref.startswith(('http:','https:','mailto:')):
        target=OUT/unquote(rel.target_ref)
        assert target.is_file(),target
        links.append(rel.target_ref)
assert len(links)==5
with (RUN/'holdout_ablation_residuals.csv').open(encoding='utf-8',newline='') as f:
    rows=list(csv.DictReader(f))
assert len(rows)==6837 and len({r['event_id'] for r in rows})==747
metrics={}
with (RUN/'holdout_ablation_metrics.csv').open(encoding='utf-8',newline='') as f:
    for row in csv.DictReader(f):
        key=row['coordinates']+'_'+row['model']
        rms=math.sqrt(sum(float(r[key])**2 for r in rows)/len(rows))
        assert abs(rms-float(row['rms_s']))<1e-12
        metrics[key]=rms
assert metrics['continuous_1d_relocated_statics']<metrics['continuous_3d_final3d_statics']
bootstrap=json.loads((RUN/'bootstrap_1d_relocalizado.json').read_text())
assert abs(metrics['continuous_3d_final3d_statics']-metrics['continuous_1d_relocated_statics']-bootstrap['continuous']['delta_rms_3d_minus_1d_s'])<1e-12
tables='\n'.join(c.text for t in D.tables for r in t.rows for c in r.cells)
for v in ['1,527','1,551','39,28','36,62']:
    assert v in tables,v
body='\n'.join(p.text for p in D.paragraphs)
assert 'Waldhauser, F.; Ellsworth, W. L.' in body
assert 'Zhang, H.; Thurber, C. H.' in body
verified_hashes=0
for p in [OUT/'fuentes/auditoria_manuscrito_revision_20261011d.json',RUN/'1d_relocated_ablation.json']:
    data=json.loads(p.read_text(encoding='utf-8'))
    for path,expected in data.get('input_sha256',data.get('source_sha256',{})).items():
        assert sha(Path(path))==expected,path
        verified_hashes+=1
old=OUT/'Manuscrito_tomografia_Vp_Colombia_UIS_revision_20261011c.docx'
assert sha(old)=='54c711d7130da94c1d83ab7d1dbe9a9d2722d7e882a4b7d3327e0231840d7a30'
with zipfile.ZipFile(old) as a,zipfile.ZipFile(DOCX) as b:
    am={n:hashlib.sha256(a.read(n)).hexdigest() for n in a.namelist() if n.startswith('word/media/')}
    bm={n:hashlib.sha256(b.read(n)).hexdigest() for n in b.namelist() if n.startswith('word/media/')}
    assert am==bm
qa=OUT/'qa_revision_20261011d_entrega'
previous=OUT/'qa_revision_20261011d_final'
changed=[p.name for p in qa.glob('page-*.png') if sha(p)!=sha(previous/p.name)]
assert sorted(changed)==['page-26.png','page-34.png','page-35.png']
audit=dict(docx_sha256=sha(DOCX),pages=35,figures=7,tables=7,local_links=links,reserve_events=747,reserve_picks=6837,recomputed_rms=metrics,bootstrap=bootstrap,verified_input_hashes=verified_hashes,previous_manuscript_unchanged=True,seven_figures_unchanged=True,visual_qa='32 pages identical to previously inspected render; pages 26, 34 and 35 reinspected after final edits',changed_pages=changed)
(OUT/'fuentes/verificacion_revision_20261011d.json').write_text(json.dumps(audit,indent=2,ensure_ascii=False),encoding='utf-8')
supp=OUT/'suplementos/diagnosticos_revision_d'
scripts=supp/'scripts';scripts.mkdir(exist_ok=True)
for name in ['ablacion_revision_20261011.py','ablacion_1d_relocalizada_20261011.py','geometria_revision_20261011.py','diagnosticos_figuras_revision_20261011.py','preparar_revision_20261011d.py','verificar_empaquetar_revision_d.py']:
    shutil.copy2(ROOT/'codigos'/name,scripts/name)
shutil.copy2(OUT/'crear_manuscrito_revision_20261011d.py',scripts/'crear_manuscrito_revision_20261011d.py')
note='''# Entrega de la revisión científica 20261011d

Descomprima todo el ZIP antes de abrir el Word: sus enlaces son relativos a suplementos/. Las vistas HTML contienen las bibliotecas necesarias para funcionar sin conexión.

Resultado principal: el control 1D con tres rondas de relocalización y correcciones de estación obtuvo RMS reservado 1,527 s; el 3D, 1,551 s. No se demuestra una ventaja predictiva del 3D frente a ese control. La respuesta a la revisión detalla los trabajos realizados y las pruebas pendientes.

Los CSV conservan la precisión completa. Los scripts científicos incluidos son copias de los ejecutados en el proyecto; requieren el entorno científico y las entradas congeladas indicadas en sus rutas y huellas. Este ZIP permite revisar resultados y código; no incluye los voluminosos campos de tiempos ni todos los datos originales necesarios para repetir la inversión. Para reproducir, use el proyecto tomografia_colombia_Q con sus directorios codigos/regional y resultados originales. Los scripts de creación de Word requieren la estructura original del proyecto y las fuentes/figuras incluidas.

Verificación: 35 páginas inspeccionadas, siete figuras conservadas, siete tablas, cinco enlaces locales comprobados, RMS de las doce combinaciones recalculados desde 6837 residuos, huellas de entradas verificadas y versión c preservada. Financiación, CRediT, conflictos e indexación de los autores requieren confirmación antes de enviar.
'''
(OUT/'LEEME_revision_20261011d.md').write_text(note,encoding='utf-8')
package=OUT/'Articulo_UIS_y_figuras_3D_revision_20261011d.zip'
files=[DOCX,OUT/'Respuesta_revision_cientifica_20261011d.md',OUT/'LEEME_revision_20261011d.md']
files.extend(p for folder in ['suplementos','figuras','fuentes'] for p in (OUT/folder).rglob('*') if p.is_file() and '__pycache__' not in p.parts)
with zipfile.ZipFile(package,'w',compression=zipfile.ZIP_DEFLATED) as z:
    for f in files:z.write(f,f.relative_to(OUT).as_posix())
with zipfile.ZipFile(package) as z:
    assert z.testzip() is None
    for target in links:assert target in z.namelist()
    assert hashlib.sha256(z.read(DOCX.name)).hexdigest()==sha(DOCX)
print(json.dumps(dict(docx=DOCX.name,zip=package.name,bytes=package.stat().st_size,files=len(files),sha256=sha(package),verification=audit),ensure_ascii=False,indent=2))
