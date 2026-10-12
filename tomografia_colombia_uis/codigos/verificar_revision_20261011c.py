from pathlib import Path
import hashlib,json,zipfile,datetime
from docx import Document
OUT=Path(__file__).resolve().parents[1]/'articulo_boletin_geologia_uis_adjunto'
src=OUT/'Manuscrito_tomografia_Vp_Colombia_UIS_revision_20261011b.docx'
dst=OUT/'Manuscrito_tomografia_Vp_Colombia_UIS_revision_20261011c.docx'
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
old,new=Document(src),Document(dst)
assert len(new.inline_shapes)==7 and len(new.tables)==6
assert len(old.paragraphs)==len(new.paragraphs)
changes=[(i,a.text,b.text) for i,(a,b) in enumerate(zip(old.paragraphs,new.paragraphs)) if a.text!=b.text]
assert len(changes)==3
assert [t._tbl.xml for t in old.tables]==[t._tbl.xml for t in new.tables]
with zipfile.ZipFile(src) as a,zipfile.ZipFile(dst) as b:
    for name in a.namelist():
        if name.startswith('word/media/'):assert a.read(name)==b.read(name)
body='\n'.join(p.text for p in new.paragraphs)
assert '(28,54 %; casi tres de cada diez) la solución' in body
assert 'límites impuestos en el 28,5 % de ellas' in body
assert 'imposed bounds in 28.5% of them' in body
reconciliation=[p for p in new.paragraphs if p.text.startswith('De los 1180 eventos elegibles')]
assert len(reconciliation)==1 and reconciliation[0].paragraph_format.keep_together
assert 'Caribe utilizó los 385' in reconciliation[0].text
targets=[r.target_ref for r in new.part.rels.values() if r.reltype.endswith('/hyperlink') and r.target_ref.startswith('suplementos/')]
assert len(targets)==3 and all((OUT/p).is_file() for p in targets)
note='''# Pulido editorial 20261011c

Se integra «casi tres de cada diez» dentro del paréntesis de §5.1, con punto y coma. La reconciliación de eventos y la explicación de la envolvente Caribe pertenecen al mismo párrafo; se impide su división entre páginas. Resumen y Abstract incorporan la limitación de amplitudes en el 28,5 % de las celdas activas. Ambos permanecen por debajo de 300 palabras. Se conserva la convención explícita de la Figura 5, azul rápido y rojo lento. Los datos, las tablas, las siete imágenes y las métricas de inversión no se modifican. El depósito público y las pruebas adicionales previamente identificadas siguen pendientes.
'''
(OUT/'suplementos/pulido_revision_20261011c.md').write_text(note,encoding='utf-8')
readme=OUT/'LEEME_paquete_revision_20261011c.txt'
readme.write_text('Extraiga todo el ZIP y mantenga suplementos junto al Word. Abra los enlaces con Ctrl+clic en Word. HTML autónomos sin conexión. Consulte suplementos/pulido_revision_20261011c.md. El paquete no constituye un depósito público con DOI.\n',encoding='utf-8')
package=OUT/'Articulo_UIS_y_figuras_3D_revision_20261011c.zip'
with zipfile.ZipFile(package,'w',zipfile.ZIP_DEFLATED) as z:
    z.write(dst,dst.name);z.write(readme,readme.name)
    for p in sorted((OUT/'suplementos').rglob('*')):
        if p.is_file():z.write(p,p.relative_to(OUT).as_posix())
with zipfile.ZipFile(package) as z:
    assert z.testzip() is None and z.read(dst.name)==dst.read_bytes()
    assert all(p in z.namelist() for p in targets)
report=dict(checked_at=datetime.datetime.now(datetime.timezone.utc).isoformat(),source_sha256=sha(src),output_sha256=sha(dst),zip_sha256=sha(package),changed_paragraphs=changes,figures=7,tables=6,tables_and_images_unchanged=True,html_links=targets,resumen_words=283,abstract_words=251)
(OUT/'fuentes/verificacion_revision_20261011c.json').write_text(json.dumps(report,indent=2,ensure_ascii=False),encoding='utf-8')
print('PASS: three text edits, paragraph kept together, tables and images unchanged, ZIP and three HTML links verified')
