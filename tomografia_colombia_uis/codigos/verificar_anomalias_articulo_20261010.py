from pathlib import Path
import hashlib, json
from docx import Document
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'articulo_boletin_geologia_uis_adjunto'
FILE = OUT / 'Manuscrito_tomografia_Vp_Colombia_UIS_con_anomalias.docx'
sha = lambda p: hashlib.sha256(Path(p).read_bytes()).hexdigest()
d = Document(FILE)
pdf = PdfReader(OUT / 'qa_anomalias/word_render.pdf')
assert len(pdf.pages) == 26 and len(d.inline_shapes) == 5 and len(d.tables) == 5
audit = json.loads((OUT / 'fuentes/auditoria_manuscrito_anomalias.json').read_text(encoding='utf-8'))
assert audit['output_sha256'] == sha(FILE)
for f, h in audit['input_sha256'].items():
    assert sha(f) == h, f
cuts = json.loads((OUT / 'fuentes/auditoria_cortes_anomalias.json').read_text(encoding='utf-8'))
for group in ['input_sha256', 'output_sha256']:
    for f, h in cuts[group].items():
        assert sha(f) == h, f
assert cuts['color_scale_percent'] == [-16, 16]
assert [x['longitude_deg'] for x in cuts['sections']] == [-77, -76, -74]
assert [x['active_section_samples'] for x in cuts['sections']] == [275, 327, 372]
text = '\n'.join(p.text for p in d.paragraphs)
assert '4.4 Cortes de anomalías de velocidad' in text
assert 'Figura 5. Cortes de anomalías porcentuales' in text
assert audit['references'] == 23
old = Document(OUT / 'Manuscrito_tomografia_Vp_Colombia_UIS_modelos_inicial_final.docx')
for a, b in zip(d.tables, old.tables):
    assert [[c.text for c in r.cells] for r in a.rows] == [[c.text for c in r.cells] for r in b.rows]
# Removing only the inserted section must reproduce every previous paragraph.
paragraphs = [p.text for p in d.paragraphs]
start = paragraphs.index('4.4 Cortes de anomalías de velocidad')
end = paragraphs.index('5 Discusión de resultados')
assert paragraphs[:start] + paragraphs[end:] == [p.text for p in old.paragraphs]
for n in range(1, 18):
    assert sha(OUT / 'qa_anomalias' / f'page-{n}.png') == sha(OUT / 'qa_modelos' / f'page-{n}.png')
for n in range(18, 27):
    assert (OUT / 'qa_anomalias' / f'page-{n}.png').is_file()
report = {'status': 'verified', 'pages': 26, 'figures': 5, 'tables': 5, 'references': 23,
          'new_render_pages_visually_inspected': list(range(18, 27)),
          'identical_to_previously_inspected_pages': list(range(1, 18)),
          'previous_paragraphs_and_tables_preserved': True,
          'figure': 5, 'section': '4.4', 'figure_page': 19,
          'output_sha256': sha(FILE), 'validation_script_sha256': sha(__file__)}
(OUT / 'fuentes/verificacion_anomalias_entrega.json').write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding='utf-8')
print(json.dumps(report, indent=2, ensure_ascii=False))
