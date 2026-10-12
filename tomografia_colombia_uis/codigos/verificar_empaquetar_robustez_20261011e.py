"""Scientific/document audit and self-contained inputs for robustness replication."""
from pathlib import Path
import json,csv,hashlib,shutil,zipfile
from urllib.parse import unquote
from docx import Document

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'articulo_boletin_geologia_uis_adjunto'
RUN=ROOT/'resultados/robustez_20261011'
ORIGINAL=ROOT/'resultados/tomografia_1d_adjunto_consistente_2026-10-10'
DOCX=OUT/'Manuscrito_tomografia_Vp_Colombia_UIS_revision_20261011e.docx'
digest=lambda f:hashlib.sha256(f.read_bytes()).hexdigest()
R=json.loads((RUN/'verified_summary.json').read_text(encoding='utf-8'))
assert R['original_inputs_unchanged']
for path,expected in json.loads((RUN/'protocol.json').read_text())['sha256'].items():assert digest(Path(path))==expected
for path,expected in json.loads((RUN/'provenance_controls.json').read_text()).items():assert digest(Path(path))==expected
frozen=json.loads((RUN/'protocol_extra.json').read_text())
for path,expected in frozen['scripts'].items():assert digest(Path(path))==expected
original_cfg=json.loads((ORIGINAL/'config.json').read_text())
for folder in [*sorted(RUN.glob('spatial_fold_*')),*sorted(RUN.glob('synthetic_*pct'))]:
    cfg=json.loads((folder/'config.json').read_text())
    for key,value in original_cfg.items():
        if key not in ['cache_directory','time_field_memory_limit_mb']:assert cfg[key]==value,(folder,key)
    history=json.loads((folder/'history.json').read_text());assert len(history)==3
    assert all(h['lsmr']['converged'] for h in history)
    assert all(h['ray_diagnostic']['failed_rows']==[] for h in history)
D=Document(DOCX);assert len(D.inline_shapes)==7 and len(D.tables)==10
text='\n'.join(p.text for p in D.paragraphs)
assert '4.7 Convergencia y validación no lineal' in text
links=[]
for rel in D.part.rels.values():
    if rel.reltype.endswith('/hyperlink') and not rel.target_ref.startswith(('http:','https:','mailto:')):
        assert (OUT/unquote(rel.target_ref)).is_file(),rel.target_ref
        links.append(rel.target_ref)
assert len(links)==9
with zipfile.ZipFile(OUT/'Manuscrito_tomografia_Vp_Colombia_UIS_revision_20261011d.docx') as a,zipfile.ZipFile(DOCX) as b:
    old={n:hashlib.sha256(a.read(n)).hexdigest() for n in a.namelist() if n.startswith('word/media/')}
    new={n:hashlib.sha256(b.read(n)).hexdigest() for n in b.namelist() if n.startswith('word/media/')}
    assert old==new
qa=OUT/'qa_revision_20261011e_entrega_final';render=json.loads((qa/'render_method.json').read_text())
audit=dict(docx_sha256=digest(DOCX),pages=len(render['pages']),figures=7,tables=10,local_links=links,lsmr_all_converged=True,failed_jacobian_rows=0,original_inputs_unchanged=True,original_seven_figures_unchanged=True,protocol_sha256=digest(RUN/'protocol.json'),scientific_summary_sha256=digest(RUN/'verified_summary.json'),visual_qa='All rendered pages inspected before delivery')
(OUT/'fuentes/verificacion_revision_20261011e.json').write_text(json.dumps(audit,indent=2,ensure_ascii=False),encoding='utf-8')
repro=OUT/'reproduccion_controles';repro.mkdir(exist_ok=True)
files=[ORIGINAL/n for n in ['config.json','domain.json','predictions.csv','vp_final.npz']]
files += list((ROOT/'codigos/regional').glob('*.py'))+[ROOT/'codigos/regional/LICENSE_REFERENCE.txt']
files += [p for p in (ROOT/'codigos/native').iterdir() if p.is_file()]
names=['robustez_20261011.py','convergencia_malla_20261011.py','sinteticos_no_lineales_20261011.py','sintetico_por_escenario_20261011.py','precalcular_sintetico_08_20261011.py','fsm_equivalente_20261011.py','resumir_robustez_20261011.py','reproducir_controles_20261011.py','nazca_campana_20261010.py','relocation.py']
files += [ROOT/'codigos'/n for n in names]
files += [ROOT/'datos/cartografia/departamentos_dane_2025.geojson']
for source in files:
    dest=repro/source.relative_to(ROOT);dest.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(source,dest)
versions=json.loads((RUN/'environment_versions.json').read_text())
(repro/'requirements.txt').write_text('\n'.join(n+'=='+v for n,v in versions.items())+'\n',encoding='utf-8')
readme='''# Reproducción de controles científicos

Esta carpeta incluye las entradas congeladas completas, el código Python, las fuentes C++ y las DLL usadas en Windows. No incluye campos de tiempos: se recalculan. El entorno registrado es Python 3.14; requirements.txt contiene las versiones usadas. En otros sistemas se deben compilar las bibliotecas nativas para su plataforma; no se validó esa ejecución.

Desde esta carpeta, ejecute con el Python científico:

    python codigos/reproducir_controles_20261011.py mesh
    python codigos/reproducir_controles_20261011.py synthetic
    python codigos/reproducir_controles_20261011.py spatial
    python codigos/reproducir_controles_20261011.py summary

El lanzador adapta únicamente el directorio de caché al directorio extraído. Todos los parámetros científicos siguen siendo los originales. Las pruebas necesitan tiempo de CPU y varios GB de espacio. La parada al punto fijo exacto es una optimización certificada opcional; la ruta anterior reproduce con el solucionador original de ocho ciclos y conserva los resultados. Los controles son retrospectivos sobre las asociaciones del catálogo original; no son una reserva obtenida de una nueva adquisición independiente.

Las columnas observed_s de los archivos synthetic_observations.csv son tiempos sintéticos. Sus identificadores y geometría proceden del catálogo real; las otras columnas de llegada originales se conservan como procedencia y no constituyen observaciones sintéticas.
'''
(repro/'LEEME.md').write_text(readme,encoding='utf-8')
shutil.copy2(RUN/'Informe_robustez_20261011.md',OUT/'Informe_robustez_20261011.md')
lines=['# Entrega científica 20261011e','','Descomprima el paquete completo y abra el Word conservando suplementos/ a su lado. Los nueve enlaces locales se comprobaron. Las vistas HTML 3D funcionan sin conexión.','','Se completaron: convergencia directa a 10, 5 y 2,5 km; dos dameros no lineales generados a 5 km e invertidos a 10 km; tres inversiones espaciales completas y sus controles 1D, con margen horizontal de 20 km. Los resultados están en el Word y en Informe_robustez_20261011.md. Realizar los ensayos no equivale a superar sus umbrales: el informe conserva los resultados negativos.','','reproduccion_controles/ contiene las entradas congeladas, fuentes y binarios para repetir los controles; sus instrucciones están en LEEME.md. Los archivos de QA de Word no forman parte de la entrega. Financiación, CRediT, conflictos e indexación de autoría aún requieren confirmación antes del envío.']
(OUT/'LEEME_revision_20261011e.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')
package=OUT/'Articulo_UIS_y_controles_robustez_revision_20261011e.zip'
files=[DOCX,OUT/'Informe_robustez_20261011.md',OUT/'LEEME_revision_20261011e.md',OUT/'crear_manuscrito_revision_20261011e.py']
files += [p for folder in ['suplementos','figuras','fuentes','reproduccion_controles'] for p in (OUT/folder).rglob('*') if p.is_file() and '__pycache__' not in p.parts]
with zipfile.ZipFile(package,'w',compression=zipfile.ZIP_DEFLATED) as z:
    for f in files:z.write(f,f.relative_to(OUT).as_posix())
with zipfile.ZipFile(package) as z:
    assert z.testzip() is None
    for target in links:assert target in z.namelist()
    assert hashlib.sha256(z.read(DOCX.name)).hexdigest()==digest(DOCX)
print(json.dumps(dict(docx=DOCX.name,archive=package.name,archive_mb=package.stat().st_size/1e6,archive_sha256=digest(package),files=len(files),audit=audit),indent=2,ensure_ascii=False))
