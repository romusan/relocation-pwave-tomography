from pathlib import Path
import json,hashlib,zipfile,datetime,xml.etree.ElementTree as ET
from docx import Document
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'articulo_boletin_geologia_uis_adjunto'
doc=OUT/'Manuscrito_tomografia_Vp_Colombia_UIS_revision_20261011b.docx'
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
d=Document(doc);texts=[p.text for p in d.paragraphs];body='\n'.join(texts)
assert texts[0].startswith('Tomografía local de Vp')
assert d.paragraphs[0].style.name=='Title'
assert len(d.inline_shapes)==7 and len(d.tables)==6
assert 'incertidumbre uniforme de 1 s' not in body
for term in ['225 restantes','150 ajustes por métrica y 300 en total','28,54 %','4.6 Evaluación de reserva por profundidad','Limitación de trazabilidad del catálogo SGC']:
    assert term in body,term
assert body.index('4.6 Evaluación')<body.index('Tabla 6.')<body.index('5 Discusión')
links=[]
for rel in d.part.rels.values():
    if rel.reltype.endswith('/hyperlink') and rel.target_ref.startswith('suplementos/'):
        assert (OUT/rel.target_ref).is_file()
        links.append(rel.target_ref)
assert len(links)==3
for p in (OUT/'suplementos').glob('*.html'):
    assert '<script src=' not in p.read_text(encoding='utf-8')
review='''# Respuesta a la revisión de la versión 20261011

La versión 20261011b mantiene los datos, la inversión, los resultados numéricos y las superficies NURBS del experimento. Se revisó el texto y se amplió la Figura 7 sin repetir la tomografía.

1. Título español: el DOCX 20261011 ya contiene el título como primer párrafo con estilo Title. La nueva versión lo conserva, hace explícita su visibilidad y se verifica en la primera página renderizada. No se duplica.
2. Selección de eventos: se contrastó asignacion_sismos.csv con el código de las superficies. Los 1180 elegibles se dividen en 795 Nazca y 385 Caribe; Nazca conserva 570 compatibles con Slab2 y excluye 225. Caribe usa sus 385. La envolvente de Caribe recorta la representación y sus anclajes, no los eventos. Se explica en §3.4.
3. Peso uniforme: se sustituye incertidumbre uniforme por divisor uniforme de residuos en §5.1.
4. Inicializaciones: las mismas 30 semillas se reutilizan en cada k y métrica; 150 ajustes por métrica, 300 en total. Se aclara que 3× de exageración gráfica no altera las distancias del ajuste.
5. Evaluación por profundidad: la Tabla 6 y su interpretación pasan a §4.6, separadas de resolución y clasificación.
6. Saturación: §5.1 cuantifica 2857/10011 = 28,54 %. Se señala que las cotas condicionan la amplitud, sin afirmar que las celdas carezcan de toda información.
7. Bisector: la Figura 7 incorpora 60, 80 y 100 km con estilos diferentes; no modifica grupos ni NURBS. El texto explica la elección de niveles alrededor del corte de 80 km.
8. Portabilidad: los enlaces pasan a la carpeta sin fecha suplementos. El paquete ZIP contiene el Word y esa carpeta contigua. La extracción completa conserva los enlaces y los HTML funcionan sin conexión.
9. Trazabilidad SGC: se separa explícitamente como limitación. No se inventa una fecha única ni se sustituye el catálogo original con una adquisición posterior. Esta revisión registra fecha y hashes de sus insumos; no los presenta como fechas de descarga histórica.

Pendiente antes del envío editorial: depositar públicamente los datos permitidos, el código y los HTML en un repositorio con identificador persistente y reemplazar enlaces locales. No se ha publicado ni emitido DOI. Permanecen pendientes las reinversiones y pruebas de resolución adicionales descritas en la revisión previa. Una descarga nueva con registro fechado constituiría otra adquisición y requeriría comprobar su equivalencia con el catálogo del experimento.
'''
(OUT/'suplementos/respuesta_revision_20261011b.md').write_text(review,encoding='utf-8')
readme='Extraiga el ZIP completo. Mantenga la carpeta suplementos junto al Word. Los tres enlaces HTML son relativos; en Word se abren mediante Ctrl+clic. Los HTML incluyen Plotly y funcionan sin conexión. El paquete es local, no un depósito público con DOI. Consulte suplementos/respuesta_revision_20261011b.md para los cambios y limitaciones pendientes.\n'
(OUT/'LEEME_paquete_revision_20261011b.txt').write_text(readme,encoding='utf-8')
inputs=[OUT/'Manuscrito_tomografia_Vp_Colombia_UIS_revision_20261011.docx',ROOT/'resultados/placas_adjunto_consistente_2026-10-10/asignacion_sismos.csv',ROOT/'codigos/placas_nazca_caribe_diagonal_20261010.py',ROOT/'codigos/suplementos_revision_20261011.py',OUT/'fuentes/metricas_articulo_revision_20261011.json']
report=dict(checked_at=datetime.datetime.now(datetime.timezone.utc).isoformat(),operation='editorial revision; not historical acquisition',input_sha256={str(p):sha(p) for p in inputs},output=str(doc),output_sha256=sha(doc),figure_count=7,table_count=6,relative_html_links=links,eligible_events=1180,nazca_candidates=795,nazca_used=570,caribe_used=385,nazca_excluded=225,saturated_cells=2857,saturated_percent=2857/10011*100)
(OUT/'suplementos/verificacion_editorial_20261011b.json').write_text(json.dumps(report,indent=2,ensure_ascii=False),encoding='utf-8')
z=OUT/'Articulo_UIS_y_figuras_3D_revision_20261011b.zip'
with zipfile.ZipFile(z,'w',zipfile.ZIP_DEFLATED) as f:
    f.write(doc,doc.name);f.write(OUT/'LEEME_paquete_revision_20261011b.txt','LEEME_paquete_revision_20261011b.txt')
    for p in sorted((OUT/'suplementos').rglob('*')):
        if p.is_file():f.write(p,p.relative_to(OUT).as_posix())
with zipfile.ZipFile(z) as f:
    assert f.testzip() is None
    assert f.read(doc.name)==doc.read_bytes()
    for target in links:assert target in f.namelist()
report['zip']=str(z);report['zip_sha256']=sha(z)
(OUT/'fuentes/verificacion_revision_20261011b.json').write_text(json.dumps(report,indent=2,ensure_ascii=False),encoding='utf-8')
print(json.dumps(report,indent=2,ensure_ascii=False))
