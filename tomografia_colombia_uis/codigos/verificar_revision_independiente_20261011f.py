"""Check independent measurements, immutable evidence and offline delivery."""
from pathlib import Path
import csv,hashlib,json,math,shutil,zipfile
from urllib.parse import unquote
from docx import Document

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'articulo_boletin_geologia_uis_adjunto'
RUN=ROOT/'resultados/revision_independiente_20261011f'
DOCX=OUT/'Manuscrito_tomografia_Vp_Colombia_UIS_revision_20261011f.docx'
sha=lambda f:hashlib.sha256(Path(f).read_bytes()).hexdigest()
readjson=lambda f:json.loads(Path(f).read_text(encoding='utf-8'))
def rows(f):
    with Path(f).open(encoding='utf-8',newline='') as stream:return list(csv.DictReader(stream))
def rms(x):return math.sqrt(sum(v*v for v in x)/len(x))
def pearson(x,y):
    mx=sum(x)/len(x);my=sum(y)/len(y)
    return sum((a-mx)*(b-my) for a,b in zip(x,y))/math.sqrt(sum((a-mx)**2 for a in x)*sum((b-my)**2 for b in y))

A=readjson(RUN/'auditoria_independiente.json')
verified=0
for path,expected in A['input_sha256'].items():
    assert sha(path)==expected,path
    copy=OUT/'reproduccion_controles'/Path(path).relative_to(ROOT)
    assert sha(copy)==expected,copy
    verified+=1
for path,expected in readjson(ROOT/'resultados/robustez_20261011/provenance_controls.json').items():
    assert sha(ROOT.parent/path)==expected,path
    verified+=1
for path,expected in readjson(ROOT/'resultados/robustez_20261011/protocol.json')['sha256'].items():
    assert sha(path)==expected,path
    verified+=1
assert sha(OUT/'Manuscrito_tomografia_Vp_Colombia_UIS_revision_20261011e.docx')=='7e743427afd91e3fb79723b19631c2e7aa518ccada00bb0ab48d5789e07af8a7'
assert sha(OUT/'Articulo_UIS_y_controles_robustez_revision_20261011e.zip')=='4eb28cf25d086bd45243ed6c2daa95a94d8428bbdd1facc15ac277b30dc30373'
errors=rows(RUN/'errores_hipocentrales_pareados.csv')
for key,result in A['synthetic_location'].items():
    subset=[r for r in errors if r['scenario']==key]
    assert len(subset)==2987 and len({r['event_id'] for r in subset})==2987
    for column,metric in [('initial_km','initial_rmse_km'),('final_3d_km','final_3d_rmse_km'),('control_1d_km','control_1d_rmse_km')]:
        assert abs(rms([float(r[column]) for r in subset])-result[metric])<1e-11
    assert abs(sum(float(r['final_3d_km'])>float(r['initial_km']) for r in subset)/len(subset)-result['fraction_worse'])<1e-14
residuals=rows(ROOT/'resultados/revision_cientifica_20261011/holdout_ablation_residuals.csv')
for r in A['depth_holdout']:
    sub=[q for q in residuals if r['depth_min_km']<=float(q['ev_prof_km']) and (float(q['ev_prof_km'])<r['depth_max_km'] or r['depth_max_km']==180 and float(q['ev_prof_km'])<=180)]
    assert len(sub)==r['picks'] and len({q['event_id'] for q in sub})==r['events']
    for key,column in [('initial_1d','continuous_1d_no_statics'),('control_1d','continuous_1d_relocated_statics'),('final_3d','continuous_3d_final3d_statics')]:
        assert abs(rms([float(q[column]) for q in sub])-r[key+'_rms_s'])<1e-12
statics=[r for r in rows(RUN/'estaticas_y_elevacion.csv') if r['elevation_m'] and int(r['train_picks'])>0 and r['station']!='CM.ACH1']
assert len(statics)==262
for key in ['1d','3d']:
    assert abs(pearson([float(r['elevation_m']) for r in statics],[float(r['static_'+key+'_s']) for r in statics])-A['statics_elevation'][key]['pearson_r'])<1e-12
reproduced=readjson(OUT/'reproduccion_controles/resultados/revision_independiente_20261011f/auditoria_independiente.json')
for key in ['synthetic_location','depth_holdout','statics_elevation','metadata_coverage']:assert reproduced[key]==A[key],key

D=Document(DOCX)
assert len(D.inline_shapes)==7 and len(D.tables)==10
body='\n'.join(p.text for p in D.paragraphs)
for invalid in ['pendiente de confirmación','queda pendiente','para completar','fuera del alcance de mínimos','En la evaluación original, los 747']:assert invalid not in body,invalid
for required in ['no recibió financiación específica','contribuyeron por igual','no existen conflictos','pendiente de mínimos cuadrados','no se presentan como estructura geológica']:assert required in body,required
table6=D.tables[5]
assert table6.cell(0,4).text=='Control 1D (s)'
for i,r in enumerate(A['depth_holdout'],1):
    for j,key in [(3,'initial_1d'),(4,'control_1d'),(5,'final_3d')]:assert table6.cell(i,j).text==f"{r[key+'_rms_s']:.4f}".replace('.',',')
links=[]
for rel in D.part.rels.values():
    if rel.reltype.endswith('/hyperlink') and not rel.target_ref.startswith(('http:','https:','mailto:')):
        name=unquote(rel.target_ref);assert (OUT/name).is_file(),name;links.append(name)
assert len(links)==12
with zipfile.ZipFile(OUT/'Manuscrito_tomografia_Vp_Colombia_UIS_revision_20261011e.docx') as old,zipfile.ZipFile(DOCX) as new:
    names=[n for n in old.namelist() if n.startswith('word/media/')]
    changed=[n for n in names if old.read(n)!=new.read(n)]
    sixth=D.inline_shapes[5]
    sixth_id=sixth._inline.graphic.graphicData.pic.blipFill.blip.embed
    assert changed==['word/'+D.part.rels[sixth_id].target_ref],changed
qa=OUT/'qa_revision_20261011f_entrega_final'
oldqa=OUT/'qa_revision_20261011f_entrega'
pages=list(qa.glob('page-*.png'));assert len(pages)==38
identical=[p.name for p in pages if sha(p)==sha(oldqa/p.name)]
audit=dict(docx_sha256=sha(DOCX),pages=38,figures=7,tables=10,local_links=links,verified_hashes=verified,
           old_e_unchanged=True,changed_figure=changed,analysis_reproduced_in_delivery_tree=True,
           visual_qa=f'All 38 first-render pages inspected; all {38-len(identical)} changed final-render pages reinspected; {len(identical)} identical pages carried forward by SHA256',identical_final_pages=identical)
(OUT/'fuentes/verificacion_revision_20261011f.json').write_text(json.dumps(audit,indent=2,ensure_ascii=False),encoding='utf-8')
shutil.copy2(RUN/'Revision_independiente.md',OUT/'Revision_independiente_20261011f.md')
supp=OUT/'suplementos/revision_independiente_20261011f'
for f in RUN.iterdir():
    if f.is_file() and f.name!='verificacion_entrega.log':shutil.copy2(f,supp/f.name)
for name in ['auditar_revision_20261011f.py','preparar_revision_independiente_20261011f.py','verificar_revision_independiente_20261011f.py']:shutil.copy2(ROOT/'codigos'/name,supp/name)
note='''# Entrega científica 20261011f

Descomprima todo el ZIP y abra el Word conservando suplementos/ a su lado. Se comprobaron los doce enlaces locales. Los HTML 3D funcionan sin conexión; la distribución local no equivale a un repositorio público persistente.

Revision_independiente_20261011f.md distingue las observaciones confirmadas, las que requieren matices y las atribuciones causales no demostradas. La versión f es un artículo diagnóstico: no presenta las anomalías como estructura validada. Incluye la comparación hipocentral con el error inicial, estáticas frente a elevación, tablas predictivas interpoladas y la figura no lineal principal. Las declaraciones están completas según confirmación de los autores.

reproduccion_controles/ incluye entradas, fuentes y DLL. Las instrucciones originales reproducen los controles e. Además, desde esa carpeta ejecute con el entorno científico registrado:

    python codigos/auditar_revision_20261011f.py

Se ejecutó esta auditoría también desde el árbol de entrega y se comprobó identidad de las cuatro secciones numéricas del JSON. No se repitieron las inversiones congeladas en esta revisión. Los campos de tiempos se recalculan al repetir los controles. Las fuentes C++ deben compilarse para otras plataformas, cuya ejecución no se verificó.

El generador Word conserva la estructura del proyecto de trabajo y usa el entorno documental. La evidencia incluye las huellas de entradas; las versiones anteriores se preservan. Se verificaron 38 páginas, siete figuras y diez tablas. Solo se sustituye la Figura 6; se mantienen las demás imágenes.

Completar la edición y ejecutar los controles no equivale a superar sus umbrales. La malla y los dameros siguen fallando los criterios documentados; no se implementó un solucionador factorizado ni se efectuó envío editorial.
'''
(OUT/'LEEME_revision_20261011f.md').write_text(note,encoding='utf-8')
repro_readme=OUT/'reproduccion_controles/LEEME.md'
text=repro_readme.read_text(encoding='utf-8')
if 'auditar_revision_20261011f.py' not in text:repro_readme.write_text(text+'\nLa auditoría posterior f se reproduce con `python codigos/auditar_revision_20261011f.py`. Sus entradas congeladas adicionales se incluyen; las métricas fueron verificadas desde este árbol de entrega.\n',encoding='utf-8')
files=[DOCX,OUT/'crear_manuscrito_revision_20261011f.py',OUT/'Revision_independiente_20261011f.md',OUT/'LEEME_revision_20261011f.md']
files.extend(p for folder in ['suplementos','figuras','fuentes','reproduccion_controles'] for p in (OUT/folder).rglob('*') if p.is_file() and '__pycache__' not in p.parts and p.name!='verificacion_entrega.log')
manifest={p.relative_to(OUT).as_posix():sha(p) for p in files}
manifest_file=OUT/'MANIFEST_revision_20261011f.json';manifest_file.write_text(json.dumps(manifest,indent=2,ensure_ascii=False),encoding='utf-8');files.append(manifest_file)
package=OUT/'Articulo_UIS_revision_independiente_20261011f.zip'
with zipfile.ZipFile(package,'w',compression=zipfile.ZIP_DEFLATED) as z:
    for f in files:z.write(f,f.relative_to(OUT).as_posix())
with zipfile.ZipFile(package) as z:
    assert z.testzip() is None
    for name,expected in manifest.items():assert hashlib.sha256(z.read(name)).hexdigest()==expected,name
    for name in links:assert name in z.namelist(),name
print(json.dumps(dict(docx=str(DOCX),zip=str(package),zip_bytes=package.stat().st_size,zip_sha256=sha(package),files=len(files),audit=audit),indent=2,ensure_ascii=False))
