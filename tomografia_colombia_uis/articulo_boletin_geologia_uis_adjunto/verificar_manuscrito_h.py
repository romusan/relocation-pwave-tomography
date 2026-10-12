"""Check delivery numbers, hyperlinks and forbidden unsupported wording."""
from pathlib import Path
import json,re,zipfile,xml.etree.ElementTree as E
from docx import Document
A=Path(__file__).resolve().parent;C=A.parent/'resultados/campana_perfil_suavidad_20261011h';path=A/'Manuscrito_tomografia_Vp_Colombia_UIS_revision_20261011h.docx'
M=json.loads((C/'manuscript_metrics.json').read_text(encoding='utf-8'));D=Document(path)
text='\n'.join(p.text for p in D.paragraphs)+'\n'+'\n'.join(cell.text for t in D.tables for row in t.rows for cell in row.cells)
bad=['por completar','campaña anterior','final nuevo','Informe de cierre','Cierre de campaña','30 pruebas','18,79','32,03','28,54','las superficies morada y naranja']
assert not [v for v in bad if v.casefold() in text.casefold()]
# «Pendiente» also denotes a regression slope; «independiente» is valid prose.
# Reject explicit editorial placeholders without rejecting either scientific term.
assert not re.search(r'(?im)^\s*pendientes?\s*[.:]?\s*$|\[\s*pendientes?\s*\]',text)
fmt=lambda x:f'{x:.3f}'.replace('.',',')
for r in M['real']:
    assert fmt(r['train']['rms_s']) in text
    assert fmt(r['holdout']['rms_s']) in text
    assert len(json.loads((C/r['arm']/'history.json').read_text()))==r['rounds']
    assert r['converged']==(r['stop']=='training_objective_plateau')
assert len(D.inline_shapes)==4
assert len(D.tables)==6
assert '324, 210–225' in text
assert 'Codex de OpenAI' in text
H=json.loads((C/'hipotesis_placas/metodologia.json').read_text(encoding='utf-8'))
assert str(H['equal_weight_cells']) in text and str(H['deep_quality_events']) in text
assert '4.7 Geometrías candidatas y división espacial' in text
for sh in H['nurbs']:
    assert f"{sh['control_cells_rms_vertical_km']:.2f}".replace('.',',') in text
    assert f"{sh['heldout_vertical_rms_km']:.2f}".replace('.',',') in text
links=[]
with zipfile.ZipFile(path) as z:
    rel=E.fromstring(z.read('word/_rels/document.xml.rels'))
    for node in rel:
        if node.attrib.get('Type','').endswith('/hyperlink'):
            target=node.attrib['Target'];links.append(target)
            if not target.startswith(('https://','http://','mailto:')):assert (A/target).is_file(),target
    xml=E.fromstring(z.read('word/document.xml'));equations=xml.findall('.//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath')
    assert len(equations)==7,len(equations)
report=dict(docx=str(path),tables=len(D.tables),main_figures=len(D.inline_shapes),editable_equations=len(equations),external_hyperlinks=links,numerical_values_match_completed_reports=True,requires_separate_visual_review=True)
(C/'document_content_qa.json').write_text(json.dumps(report,indent=2,ensure_ascii=False),encoding='utf-8');print(json.dumps(report,indent=2,ensure_ascii=False))
